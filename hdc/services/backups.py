"""HDC services.backups — moved verbatim from hdc_erp.py.

See MODULARIZATION_PLAN.md for the module map.
"""

import os
import shutil
import sqlite3
import tempfile
import zipfile
from datetime import datetime

import openpyxl

from hdc.config import BASE_DIR, get_runtime_settings
from hdc.utils.dates import PKT_ZONE, _pkt_now_naive
from hdc.utils.format import _quote_ident, _safe_sheet_name

def _backup_filename():
    return f"hdc_backup_{_pkt_now_naive().strftime('%Y%m%d_%H%M%S')}.zip"


def _create_backup_xlsx(dst_path):
    settings = get_runtime_settings()
    db_path = settings.db_path
    wb = openpyxl.Workbook()
    if wb.active:
        wb.remove(wb.active)

    # Metadata sheet
    ws_meta = wb.create_sheet(title='README')
    ws_meta.append(['Generated At (PKT)', _pkt_now_naive().strftime('%Y-%m-%d %H:%M:%S')])
    ws_meta.append(['Database Path', db_path])
    ws_meta.append(['Note', 'This file is a tabular export backup (not a direct restore source).'])

    con = sqlite3.connect(db_path)
    cur = con.cursor()
    cur.execute("""
        SELECT name
        FROM sqlite_master
        WHERE type='table' AND name NOT LIKE 'sqlite_%'
        ORDER BY name
    """)
    table_names = [r[0] for r in cur.fetchall()]

    used_sheet_names = {'README'}
    for tname in table_names:
        sheet_name = _safe_sheet_name(tname, fallback='Table')
        base_name = sheet_name
        suffix = 2
        while sheet_name in used_sheet_names:
            tail = f'_{suffix}'
            sheet_name = (base_name[: max(1, 31 - len(tail))] + tail)
            suffix += 1
        used_sheet_names.add(sheet_name)

        ws = wb.create_sheet(title=sheet_name)
        cur.execute(f"PRAGMA table_info({_quote_ident(tname)})")
        cols = [r[1] for r in cur.fetchall()]
        if not cols:
            ws.append(['(no columns detected)'])
            continue
        ws.append(cols)
        ws.freeze_panes = 'A2'

        col_csv = ', '.join(_quote_ident(c) for c in cols)
        q = f"SELECT {col_csv} FROM {_quote_ident(tname)}"
        for row in cur.execute(q):
            out = []
            for v in row:
                if isinstance(v, (bytes, bytearray)):
                    out.append(v.hex())
                else:
                    out.append(v)
            ws.append(out)

    con.close()
    wb.save(dst_path)


def _create_backup_zip(dst_path):
    settings = get_runtime_settings()
    fd, xlsx_path = tempfile.mkstemp(prefix='hdc_data_export_', suffix='.xlsx', dir=settings.instance_dir)
    os.close(fd)
    db_fd, db_snapshot_path = tempfile.mkstemp(prefix='hdc_db_snapshot_', suffix='.db', dir=settings.instance_dir)
    os.close(db_fd)
    try:
        _create_backup_xlsx(xlsx_path)
        # Take a consistent SQLite snapshot so WAL data is included in backup DB.
        src = sqlite3.connect(settings.db_path)
        try:
            dst = sqlite3.connect(db_snapshot_path)
            try:
                src.backup(dst)
                dst.commit()
                # Sanity-check snapshot integrity and presence of critical tables.
                cur = dst.cursor()
                for t in ('hdc_user', 'hdc_project', 'hdc_account', 'hdc_account_txn'):
                    cur.execute("SELECT name FROM sqlite_master WHERE type='table' AND name=?", (t,))
                    if cur.fetchone() is None:
                        raise RuntimeError(f'Backup snapshot missing required table: {t}')
                cur.close()
            finally:
                dst.close()
        finally:
            src.close()
        with zipfile.ZipFile(dst_path, 'w', compression=zipfile.ZIP_DEFLATED) as zf:
            if os.path.exists(db_snapshot_path):
                zf.write(db_snapshot_path, arcname='hdc_erp.db')
            if os.path.exists(settings.estimation_store):
                zf.write(settings.estimation_store, arcname='project_estimations.json')
            if os.path.exists(xlsx_path):
                zf.write(xlsx_path, arcname='hdc_data_export.xlsx')
    finally:
        try:
            if os.path.exists(db_snapshot_path):
                os.remove(db_snapshot_path)
        except Exception:
            pass
        try:
            if os.path.exists(xlsx_path):
                os.remove(xlsx_path)
        except Exception:
            pass


def _list_backups():
    settings = get_runtime_settings()
    rows = []
    for name in os.listdir(settings.backup_dir):
        if not name.lower().endswith('.zip'):
            continue
        p = os.path.join(settings.backup_dir, name)
        try:
            st = os.stat(p)
            # Always display backup file times in Pakistan Standard Time.
            modified_pkt = datetime.fromtimestamp(st.st_mtime, PKT_ZONE)
            rows.append({
                'name': name,
                'size': st.st_size,
                'modified': modified_pkt
            })
        except OSError:
            continue
    rows.sort(key=lambda r: r['modified'], reverse=True)
    return rows


def _prune_saved_backups(keep_latest=10):
    settings = get_runtime_settings()
    try:
        keep = int(keep_latest or 0)
    except Exception:
        keep = 10
    keep = max(0, min(200, keep))
    rows = _list_backups()
    deleted = 0
    errors = 0
    for r in rows[keep:]:
        name = (r.get('name') or '').strip()
        if not name:
            continue
        p = os.path.join(settings.backup_dir, name)
        try:
            if os.path.exists(p):
                os.remove(p)
                deleted += 1
        except Exception:
            errors += 1
    return {'keep': keep, 'deleted': deleted, 'errors': errors}


def _cleanup_backup_temp_artifacts():
    settings = get_runtime_settings()
    removed_files = 0
    removed_dirs = 0
    # Temporary snapshot/export files that may remain after interrupted backups.
    try:
        for name in os.listdir(settings.instance_dir):
            p = os.path.join(settings.instance_dir, name)
            if not os.path.isfile(p):
                continue
            if (name.startswith('hdc_db_snapshot_') and name.endswith('.db')) or \
               (name.startswith('hdc_data_export_') and name.endswith('.xlsx')):
                try:
                    os.remove(p)
                    removed_files += 1
                except Exception:
                    pass
    except Exception:
        pass
    # Temporary extraction folders that may remain after interrupted restore.
    try:
        for name in os.listdir(BASE_DIR):
            p = os.path.join(BASE_DIR, name)
            if (not os.path.isdir(p)) or (not name.startswith('_restore_extract_')):
                continue
            try:
                shutil.rmtree(p, ignore_errors=True)
                removed_dirs += 1
            except Exception:
                pass
    except Exception:
        pass
    return {'files': removed_files, 'dirs': removed_dirs}
