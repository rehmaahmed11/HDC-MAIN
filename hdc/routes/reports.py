"""HDC routes: Reports and CSV/XLSX/print exports.

Moved verbatim from hdc_erp.py; each handler keeps its
original @app.route decorator and endpoint name.
"""

import csv
import io
from datetime import datetime, timedelta

import openpyxl
from flask import Response, render_template, request
from flask_login import login_required
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import func

from hdc.extensions import db
from hdc.models.accounts import Alert, Expense, ExpenseCategory, OwnerPayment
from hdc.models.materials import Material, MaterialUsage, Purchase, UsageLogV2
from hdc.models.office import OfficeExpense
from hdc.models.projects import Project, Stage
from hdc.models.subcontract import SubcontractLabourAttendance, SubcontractPayment, SubcontractTeamAttendance, Subcontractor
from hdc.models.workforce import LabourLedger, PayrollItem, PayrollRun, TimeEntry, Worker
from hdc.services.aggregation import _apply_aggregated_project_costs, _apply_aggregated_stage_costs
from hdc.services.ledger import _office_expense_total
from hdc.services.lookups import _trade_options
from hdc.services.purchase import _material_stock_map
from hdc.services.reporting import _project_report_data
from hdc.utils.dates import _pkt_now, _pkt_today
from hdc.utils.format import _parse_date

def register(app):
    """Register Reports and CSV/XLSX/print exports."""
    # â”€â”€ Reports â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€
    @app.route('/hdc/reports')
    @login_required
    def hdc_reports():
        projects = Project.query.all()
        _apply_aggregated_project_costs(projects)
        workers = Worker.query.order_by(Worker.name).all()
        trade_options = _trade_options()
        materials = Material.query.order_by(Material.name).all()
        stages = Stage.query.order_by(Stage.id).all()
        pid = request.args.get('project_id', type=int)
        selected = Project.query.get(pid) if pid else None
        section = request.args.get('section', 'project')

        worker_project_id = request.args.get('worker_project_id', type=int)
        worker_stage_id = request.args.get('worker_stage_id', type=int)
        worker_worker_id = request.args.get('worker_worker_id', type=int)
        worker_trade = (request.args.get('worker_trade') or '').strip()
        worker_date_from = request.args.get('worker_date_from')
        worker_date_to = request.args.get('worker_date_to')
        # Worker report safety: when project is "All", ignore stale stage filter values.
        if not worker_project_id:
            worker_stage_id = None
        elif worker_stage_id:
            st = Stage.query.get(worker_stage_id)
            if (not st) or int(st.project_id or 0) != int(worker_project_id):
                worker_stage_id = None

        material_project_id = request.args.get('material_project_id', type=int)
        material_stage_id = request.args.get('material_stage_id', type=int)
        material_id = request.args.get('material_id', type=int)
        material_date_from = request.args.get('material_date_from')
        material_date_to = request.args.get('material_date_to')

        time_project_id = request.args.get('time_project_id', type=int)
        time_stage_id = request.args.get('time_stage_id', type=int)
        time_worker_id = request.args.get('time_worker_id', type=int)
        time_worker_name = (request.args.get('time_worker_name') or '').strip()
        time_trade = (request.args.get('time_trade') or '').strip()
        time_date_from = request.args.get('time_date_from')
        time_date_to = request.args.get('time_date_to')

        stage_project_id = request.args.get('stage_project_id', type=int)
        stage_id = request.args.get('stage_id', type=int)
        expense_project_id = request.args.get('expense_project_id', type=int)
        expense_stage_id = request.args.get('expense_stage_id', type=int)
        expense_category = (request.args.get('expense_category') or '').strip()
        expense_date_from = request.args.get('expense_date_from')
        expense_date_to = request.args.get('expense_date_to')

        stage_rows = []
        if selected:
            selected_stages = Stage.query.filter_by(project_id=pid).order_by(Stage.id).all()
            _apply_aggregated_stage_costs(selected_stages)
            for s in selected_stages:
                stage_rows.append({
                    'name': s.name,
                    'status': s.status,
                    'contract_value': s.contract_value,
                    'labour': s.stage_labour_cost,
                    'materials': s.stage_material_cost,
                    'expenses': s.stage_expense_cost,
                    'subcontracts': s.stage_subcontract_cost,
                    'total_cost': s.stage_total_cost,
                    'profit': s.stage_profit,
                })

        # Worker report (summary for all workers, detail for single worker)
        def _apply_worker_filters(q):
            if worker_project_id:
                q = q.filter(TimeEntry.project_id == worker_project_id)
            if worker_stage_id:
                q = q.filter(TimeEntry.stage_id == worker_stage_id)
            if worker_worker_id:
                q = q.filter(TimeEntry.worker_id == worker_worker_id)
            if worker_trade:
                q = q.filter(Worker.role_type.ilike(f"%{worker_trade}%"))
            if worker_date_from:
                q = q.filter(TimeEntry.check_in >= datetime.strptime(worker_date_from, '%Y-%m-%d'))
            if worker_date_to:
                q = q.filter(TimeEntry.check_in < datetime.strptime(worker_date_to, '%Y-%m-%d') + timedelta(days=1))
            return q

        worker_summary_rows = []
        worker_detail_rows = []
        worker_selected = Worker.query.get(worker_worker_id) if worker_worker_id else None
        worker_records = []

        if worker_worker_id:
            wdq = (db.session.query(
                    func.date(TimeEntry.check_in).label('work_date'),
                    Worker.name.label('worker_name'),
                    Project.name.label('project_name'),
                    Stage.name.label('stage_name'),
                    func.coalesce(func.sum(TimeEntry.hours), 0.0).label('total_hours'),
                    func.coalesce(func.sum(TimeEntry.wage_calculated), 0.0).label('total_wage')
                  )
                  .join(Worker, TimeEntry.worker_id == Worker.id)
                  .join(Project, TimeEntry.project_id == Project.id)
                  .outerjoin(Stage, TimeEntry.stage_id == Stage.id)
                  .filter(TimeEntry.is_void == False))
            wdq = _apply_worker_filters(wdq)
            worker_detail_rows = (wdq.group_by(
                                    func.date(TimeEntry.check_in),
                                    Worker.name,
                                    Project.name,
                                    Stage.name
                                  )
                                  .order_by(func.date(TimeEntry.check_in).desc(), Project.name.asc(), Stage.name.asc())
                                  .all())
            worker_grand_days = len({str(r.work_date) for r in worker_detail_rows if r.work_date})
            worker_grand_hours = sum(float(r.total_hours or 0) for r in worker_detail_rows)
            worker_grand_wage = sum(float(r.total_wage or 0) for r in worker_detail_rows)
        else:
            wsq = (db.session.query(
                    Worker.id.label('worker_id'),
                    Worker.name.label('worker_name'),
                    Worker.role_type.label('worker_trade'),
                    func.count(func.distinct(func.date(TimeEntry.check_in))).label('days_worked'),
                    func.coalesce(func.sum(TimeEntry.hours), 0.0).label('total_hours'),
                    func.coalesce(func.sum(TimeEntry.wage_calculated), 0.0).label('total_wage')
                  )
                  .join(Worker, TimeEntry.worker_id == Worker.id)
                  .join(Project, TimeEntry.project_id == Project.id)
                  .filter(TimeEntry.is_void == False))
            wsq = _apply_worker_filters(wsq)
            worker_summary_rows = (wsq.group_by(Worker.id, Worker.name, Worker.role_type)
                                   .order_by(Worker.name.asc())
                                   .all())
            worker_grand_days = sum(int(r.days_worked or 0) for r in worker_summary_rows)
            worker_grand_hours = sum(float(r.total_hours or 0) for r in worker_summary_rows)
            worker_grand_wage = sum(float(r.total_wage or 0) for r in worker_summary_rows)

        worker_total_wage = worker_grand_wage

        # Material report (purchases)
        mq = (db.session.query(Purchase, Material, Project)
              .join(Material, Purchase.material_id == Material.id)
              .join(Project, Purchase.project_id == Project.id))
        if material_project_id:
            mq = mq.filter(Purchase.project_id == material_project_id)
        if material_stage_id:
            mq = mq.filter(Purchase.stage_id == material_stage_id)
        if material_id:
            mq = mq.filter(Purchase.material_id == material_id)
        if material_date_from:
            mq = mq.filter(Purchase.date >= _parse_date(material_date_from))
        if material_date_to:
            mq = mq.filter(Purchase.date <= _parse_date(material_date_to))
        material_records = mq.order_by(Purchase.date.desc()).all()
        material_total_cost = sum(float(pur.total or 0) for pur, _, _ in material_records)

        # Time tracking report
        tq = (db.session.query(TimeEntry, Worker, Project)
              .join(Worker, TimeEntry.worker_id == Worker.id)
              .join(Project, TimeEntry.project_id == Project.id)
              .filter(TimeEntry.is_void == False))
        if time_project_id:
            tq = tq.filter(TimeEntry.project_id == time_project_id)
        if time_stage_id:
            tq = tq.filter(TimeEntry.stage_id == time_stage_id)
        if time_worker_id:
            tq = tq.filter(TimeEntry.worker_id == time_worker_id)
        if time_worker_name:
            tq = tq.filter(Worker.name.ilike(f"%{time_worker_name}%"))
        if time_trade:
            tq = tq.filter(Worker.role_type.ilike(f"%{time_trade}%"))
        if time_date_from:
            tq = tq.filter(TimeEntry.check_in >= datetime.strptime(time_date_from, '%Y-%m-%d'))
        if time_date_to:
            tq = tq.filter(TimeEntry.check_in <= datetime.strptime(time_date_to, '%Y-%m-%d') + timedelta(days=1))
        time_records = tq.order_by(TimeEntry.check_in.desc()).all()
        time_total_hours = sum(float(t.hours or 0) for t, _, _ in time_records)
        time_total_wage = sum(float(t.wage_calculated or 0) for t, _, _ in time_records)

        # Stage cost report
        stage_cost_rows = []
        stage_q = Stage.query
        if stage_project_id:
            stage_q = stage_q.filter(Stage.project_id == stage_project_id)
        if stage_id:
            stage_q = stage_q.filter(Stage.id == stage_id)
        stage_source = stage_q.all()
        _apply_aggregated_stage_costs(stage_source)
        for s in stage_source:
            stage_cost_rows.append({
                'project': s.project.name if s.project else '',
                'stage': s.name,
                'estimated': s.estimated_cost or 0,
                'actual': s.actual_cost or 0,
                'progress': s.progress or 0,
            })

        # Expense breakdown
        eq = Expense.query.join(ExpenseCategory, Expense.category_id == ExpenseCategory.id)
        eq = eq.filter(Expense.is_void == False)
        if expense_project_id:
            eq = eq.filter(Expense.project_id == expense_project_id)
        if expense_stage_id:
            eq = eq.filter(Expense.stage_id == expense_stage_id)
        if expense_category:
            eq = eq.filter(ExpenseCategory.name.ilike(f"%{expense_category}%"))
        if expense_date_from:
            eq = eq.filter(Expense.date >= _parse_date(expense_date_from))
        if expense_date_to:
            eq = eq.filter(Expense.date <= _parse_date(expense_date_to))
        expense_records = eq.order_by(Expense.date.desc()).all()
        expense_total = sum(float(e.amount or 0) for e in expense_records)

        return render_template('reports/reports.html',
            projects=projects, workers=workers, materials=materials, stages=stages,
            trade_options=trade_options,
            selected=selected, pid=pid, section=section,
            stage_rows=stage_rows,
            worker_records=worker_records, worker_project_id=worker_project_id,
            worker_stage_id=worker_stage_id, worker_worker_id=worker_worker_id, worker_trade=worker_trade,
            worker_date_from=worker_date_from, worker_date_to=worker_date_to,
            worker_total_wage=worker_total_wage,
            worker_selected=worker_selected,
            worker_summary_rows=worker_summary_rows,
            worker_detail_rows=worker_detail_rows,
            worker_grand_days=worker_grand_days,
            worker_grand_hours=worker_grand_hours,
            worker_grand_wage=worker_grand_wage,
            material_records=material_records, material_project_id=material_project_id,
            material_stage_id=material_stage_id, material_id=material_id,
            material_date_from=material_date_from, material_date_to=material_date_to,
            material_total_cost=material_total_cost,
            time_records=time_records, time_project_id=time_project_id,
            time_stage_id=time_stage_id, time_worker_id=time_worker_id,
            time_worker_name=time_worker_name, time_trade=time_trade,
            time_date_from=time_date_from, time_date_to=time_date_to,
            time_total_hours=time_total_hours, time_total_wage=time_total_wage,
            stage_cost_rows=stage_cost_rows, stage_project_id=stage_project_id, stage_id=stage_id,
            expense_records=expense_records, expense_project_id=expense_project_id,
            expense_stage_id=expense_stage_id, expense_category=expense_category,
            expense_date_from=expense_date_from, expense_date_to=expense_date_to,
            expense_total=expense_total)


    @app.route('/hdc/reports/glance')
    @login_required
    def hdc_reports_glance():
        today = _pkt_today()
        start_30 = today - timedelta(days=29)
        start_today_dt = datetime.combine(today, datetime.min.time())
        end_today_dt = start_today_dt + timedelta(days=1)

        projects = Project.query.order_by(Project.id.desc()).all()
        _apply_aggregated_project_costs(projects)
        stages = Stage.query.order_by(Stage.id.desc()).all()
        _apply_aggregated_stage_costs(stages)
        subcontractors = Subcontractor.query.order_by(Subcontractor.id.desc()).all()
        materials = Material.query.order_by(Material.name.asc(), Material.id.asc()).all()

        total_contract = float(sum(float(p.owner_contract_value or 0.0) for p in projects))
        total_cost = float(sum(float(p.total_cost or 0.0) for p in projects))
        total_received = float(sum(float(p.total_received or 0.0) for p in projects))
        office_expense_total = _office_expense_total()
        total_profit = float(sum(float(p.net_profit or 0.0) for p in projects)) - office_expense_total
        total_receivable = float(sum(float(p.remaining_receivable or 0.0) for p in projects))
        collection_pct = (total_received / total_contract * 100.0) if total_contract > 0 else 0.0

        active_projects = sum(1 for p in projects if (p.status or '').strip().lower() == 'active')
        completed_stages = sum(1 for s in stages if (s.status or '').strip().lower() in ('completed', 'complete'))
        delayed_stages = [
            s for s in stages
            if s.end_date and s.end_date < today and (s.status or '').strip().lower() not in ('completed', 'complete')
        ]
        delayed_stage_rows = sorted(
            delayed_stages,
            key=lambda s: (s.end_date or today, s.id or 0)
        )[:10]
        stage_over_budget_count = sum(
            1 for s in stages
            if float(s.estimated_cost or 0.0) > 0 and float(s.actual_cost or 0.0) > float(s.estimated_cost or 0.0)
        )
        project_budget_overrun_rows = []
        for p in projects:
            budget = float(p.budget_total or 0.0)
            actual = float(p.total_cost or 0.0)
            if budget > 0 and actual > budget:
                project_budget_overrun_rows.append({
                    'project': p,
                    'budget': budget,
                    'actual': actual,
                    'overrun': actual - budget
                })
        project_budget_overrun_rows = sorted(project_budget_overrun_rows, key=lambda r: r['overrun'], reverse=True)[:10]

        expense_30 = float(db.session.query(func.coalesce(func.sum(Expense.amount), 0.0))
                           .filter(Expense.date >= start_30, Expense.date <= today, Expense.is_void == False)
                           .scalar() or 0.0)
        sub_pay_30 = float(db.session.query(func.coalesce(func.sum(SubcontractPayment.amount), 0.0))
                           .filter(SubcontractPayment.date >= start_30, SubcontractPayment.date <= today, SubcontractPayment.is_void == False)
                           .scalar() or 0.0)
        labour_cash_30 = float(db.session.query(func.coalesce(func.sum(LabourLedger.amount), 0.0))
                               .filter(
                                   LabourLedger.is_void == False,
                                   LabourLedger.entry_type.in_(['payment', 'advance']),
                                   LabourLedger.date >= start_30,
                                   LabourLedger.date <= today
                               )
                               .scalar() or 0.0)
        material_net_30 = float(db.session.query(func.coalesce(func.sum(Purchase.total), 0.0))
                                .filter(Purchase.date >= start_30, Purchase.date <= today)
                                .scalar() or 0.0)
        material_usage_30_old = float(db.session.query(func.coalesce(func.sum(MaterialUsage.total), 0.0))
                                      .filter(MaterialUsage.used_at >= start_30, MaterialUsage.used_at <= today)
                                      .scalar() or 0.0)
        material_usage_30_v2 = float(db.session.query(func.coalesce(func.sum(UsageLogV2.cost), 0.0))
                                     .filter(UsageLogV2.is_void == False, UsageLogV2.date >= start_30, UsageLogV2.date <= today)
                                     .scalar() or 0.0)
        material_usage_30 = material_usage_30_old + material_usage_30_v2
        received_30 = float(db.session.query(func.coalesce(func.sum(OwnerPayment.amount), 0.0))
                            .filter(OwnerPayment.date >= start_30, OwnerPayment.date <= today, OwnerPayment.is_void == False)
                            .scalar() or 0.0)
        office_expense_30 = float(db.session.query(func.coalesce(func.sum(OfficeExpense.amount), 0.0))
                                  .filter(OfficeExpense.is_void == False, OfficeExpense.date >= start_30, OfficeExpense.date <= today)
                                  .scalar() or 0.0)
        cash_out_30 = expense_30 + office_expense_30 + sub_pay_30 + labour_cash_30 + material_net_30
        cash_delta_30 = received_30 - cash_out_30

        active_workers = int(Worker.query.filter(Worker.active_status == True).count())
        workers_present_today = int(db.session.query(func.count(func.distinct(TimeEntry.worker_id)))
                                    .filter(
                                        TimeEntry.is_void == False,
                                        TimeEntry.check_in >= start_today_dt,
                                        TimeEntry.check_in < end_today_dt
                                    )
                                    .scalar() or 0)
        work_hours_today = float(db.session.query(func.coalesce(func.sum(TimeEntry.hours), 0.0))
                                 .filter(
                                     TimeEntry.is_void == False,
                                     TimeEntry.check_in >= start_today_dt,
                                     TimeEntry.check_in < end_today_dt
                                 )
                                 .scalar() or 0.0)
        wages_today = float(db.session.query(func.coalesce(func.sum(TimeEntry.wage_calculated), 0.0))
                            .filter(
                                TimeEntry.is_void == False,
                                TimeEntry.check_in >= start_today_dt,
                                TimeEntry.check_in < end_today_dt
                            )
                            .scalar() or 0.0)

        sub_labour_cost_map = {
            int(sid): float(total or 0.0)
            for sid, total in (
                db.session.query(
                    SubcontractLabourAttendance.subcontractor_id,
                    func.coalesce(func.sum(SubcontractLabourAttendance.total_labour_paid), 0.0)
                )
                .group_by(SubcontractLabourAttendance.subcontractor_id)
                .all()
            ) if sid
        }
        # Simple crew summaries (trade x days x workers x rate) feed the same
        # labour cost so reports match the subcontractor attendance page.
        for sid, total in (
            db.session.query(
                SubcontractTeamAttendance.subcontractor_id,
                func.coalesce(func.sum(SubcontractTeamAttendance.total_amount), 0.0)
            )
            .group_by(SubcontractTeamAttendance.subcontractor_id)
            .all()
        ):
            if sid:
                sub_labour_cost_map[int(sid)] = sub_labour_cost_map.get(int(sid), 0.0) + float(total or 0.0)
        subcontract_total_payable = float(sum(float(s.payable_amount or 0.0) for s in subcontractors))
        subcontract_total_cleared = float(sum(float(s.total_cleared or 0.0) for s in subcontractors))
        subcontract_total_balance = float(sum(float(s.payable_balance or 0.0) for s in subcontractors))
        subcontract_watch_rows = []
        subcontract_loss_count = 0
        for s in subcontractors:
            labour_cost = float(sub_labour_cost_map.get(int(s.id), 0.0))
            pnl = float(s.contract_value or 0.0) - labour_cost
            if pnl < 0:
                subcontract_loss_count += 1
            if float(s.payable_balance or 0.0) > 0 or pnl < 0:
                subcontract_watch_rows.append({
                    'sub': s,
                    'labour_cost': labour_cost,
                    'pnl': pnl
                })
        subcontract_watch_rows = sorted(
            subcontract_watch_rows,
            key=lambda r: (float(r['pnl']), -float(r['sub'].payable_balance or 0.0))
        )[:10]

        stock_map = _material_stock_map()
        low_stock_rows = []
        for m in materials:
            row = stock_map.get(m.id, {})
            low_stock_rows.append({
                'material': m,
                'purchased': float(row.get('purchased', 0.0) or 0.0),
                'used': float(row.get('used', 0.0) or 0.0),
                'remaining': float(row.get('remaining', 0.0) or 0.0)
            })
        low_stock_rows = sorted(low_stock_rows, key=lambda r: float(r['remaining']))[:10]
        out_of_stock_count = sum(1 for r in low_stock_rows if float(r['remaining'] or 0.0) <= 0)

        risk_project_rows = sorted(
            [
                {
                    'project': p,
                    'profit': float(p.net_profit or 0.0),
                    'remaining': float(p.remaining_receivable or 0.0),
                    'cost_ratio': ((float(p.total_cost or 0.0) / float(p.owner_contract_value or 1.0)) * 100.0) if float(p.owner_contract_value or 0.0) > 0 else 0.0
                }
                for p in projects
                if float(p.net_profit or 0.0) < 0 or float(p.remaining_receivable or 0.0) > 0
            ],
            key=lambda r: (float(r['profit']), -float(r['remaining']))
        )[:10]

        open_alerts = (Alert.query
                       .filter(Alert.resolved == False)
                       .order_by(Alert.created_at.desc(), Alert.id.desc())
                       .limit(12)
                       .all())

        latest_payroll = PayrollRun.query.order_by(PayrollRun.run_date.desc(), PayrollRun.id.desc()).first()
        latest_payroll_workers = 0
        if latest_payroll:
            latest_payroll_workers = int(db.session.query(func.count(PayrollItem.id))
                                         .filter(PayrollItem.run_id == latest_payroll.id)
                                         .scalar() or 0)

        return render_template('reports/reports_glance.html',
            today=today,
            start_30=start_30,
            active_projects=active_projects,
            project_count=len(projects),
            stage_count=len(stages),
            completed_stages=completed_stages,
            delayed_stages_count=len(delayed_stages),
            stage_over_budget_count=stage_over_budget_count,
            total_contract=total_contract,
            total_cost=total_cost,
            total_received=total_received,
            total_profit=total_profit,
            office_expense_total=office_expense_total,
            total_receivable=total_receivable,
            collection_pct=collection_pct,
            received_30=received_30,
            cash_out_30=cash_out_30,
            cash_delta_30=cash_delta_30,
            expense_30=expense_30,
            office_expense_30=office_expense_30,
            sub_pay_30=sub_pay_30,
            labour_cash_30=labour_cash_30,
            material_net_30=material_net_30,
            material_usage_30=material_usage_30,
            active_workers=active_workers,
            workers_present_today=workers_present_today,
            work_hours_today=work_hours_today,
            wages_today=wages_today,
            subcontract_count=len(subcontractors),
            subcontract_total_payable=subcontract_total_payable,
            subcontract_total_cleared=subcontract_total_cleared,
            subcontract_total_balance=subcontract_total_balance,
            subcontract_loss_count=subcontract_loss_count,
            out_of_stock_count=out_of_stock_count,
            risk_project_rows=risk_project_rows,
            delayed_stage_rows=delayed_stage_rows,
            project_budget_overrun_rows=project_budget_overrun_rows,
            subcontract_watch_rows=subcontract_watch_rows,
            low_stock_rows=low_stock_rows,
            open_alerts=open_alerts,
            latest_payroll=latest_payroll,
            latest_payroll_workers=latest_payroll_workers
        )


    @app.route('/hdc/reports/export/profitability')
    @login_required
    def hdc_export_profitability():
        out = io.StringIO()
        w   = csv.writer(out)
        w.writerow(['Code','Name','Client','Contract Value','Stage Value','Total Received',
                    'Labour Cost','Material Cost','Expense Cost','Subcontract Cost',
                    'Total Cost','Gross Margin','Net Profit','Remaining'])
        projects = Project.query.all()
        _apply_aggregated_project_costs(projects)
        for p in projects:
            w.writerow([p.project_code, p.name, p.client or '',
                        p.owner_contract_value, p.stage_contract_value,
                        p.total_received, p.total_labour_cost, p.total_material_cost,
                        p.total_expense_cost, p.total_subcontract_cost,
                        p.total_cost, p.gross_margin, p.net_profit, p.remaining_receivable])
        out.seek(0)
        return Response(out.getvalue(), mimetype='text/csv',
            headers={'Content-Disposition': 'attachment; filename=profitability_report.csv'})


    @app.route('/hdc/reports/export/salary')
    @login_required
    def hdc_export_salary():
        out = io.StringIO()
        w   = csv.writer(out)
        w.writerow(['Date','Worker Code','Worker Name','Role','Project','Stage','Hours','Overtime','Wage'])
        for t, wk, p in (db.session.query(TimeEntry, Worker, Project)
                         .join(Worker, TimeEntry.worker_id == Worker.id)
                         .join(Project, TimeEntry.project_id == Project.id)
                         .filter(TimeEntry.is_void == False)
                         .order_by(TimeEntry.check_in).all()):
            w.writerow([t.check_in.date(), wk.worker_code, wk.name, wk.role_type or '', p.name,
                        t.stage.name if t.stage_id else '',
                        t.hours, t.overtime, t.wage_calculated])
        out.seek(0)
        return Response(out.getvalue(), mimetype='text/csv',
            headers={'Content-Disposition': 'attachment; filename=salary_report.csv'})


    @app.route('/hdc/reports/export/materials')
    @login_required
    def hdc_export_materials():
        out = io.StringIO()
        w   = csv.writer(out)
        w.writerow(['Date','Type','Project','Stage','Material','Unit','Qty','Rate','Total','Supplier','Return Ref','Reason','Approved By'])
        for pur, mat, p in (db.session.query(Purchase, Material, Project)
                            .join(Material, Purchase.material_id == Material.id)
                            .join(Project, Purchase.project_id == Project.id)
                            .order_by(Purchase.date).all()):
            w.writerow([pur.date, (pur.entry_type or 'purchase'), p.name, pur.stage.name if pur.stage_id else '',
                        mat.name, mat.unit, pur.qty, pur.rate, pur.total,
                        pur.supplier_name or '', pur.return_ref or '', pur.return_reason or '', pur.approved_by or ''])
        out.seek(0)
        return Response(out.getvalue(), mimetype='text/csv',
            headers={'Content-Disposition': 'attachment; filename=materials_report.csv'})


    @app.route('/hdc/reports/project/<int:pid>/csv')
    @login_required
    def hdc_project_report_csv(pid):
        d   = _project_report_data(pid)
        p   = d['project']
        out = io.StringIO()
        w   = csv.writer(out)

        w.writerow(['HDC ERP â€“ Project Report', p.name])
        w.writerow(['Client', p.client or ''])
        w.writerow(['Location', p.location or ''])
        w.writerow(['Contract Type', p.contract_type])
        w.writerow(['Contract Value (Rs)', f"{d['contract_value']:,.2f}"])
        w.writerow(['Total Cost (Rs)', f"{d['total_cost']:,.2f}"])
        w.writerow(['Profit (Rs)', f"{d['profit']:,.2f}"])
        w.writerow([])

        # Stage Summary
        w.writerow(['== Stage Summary =='])
        w.writerow(['Stage', 'Status', 'Stage Value (Rs)', 'Labour (Rs)',
                    'Materials (Rs)', 'Expenses (Rs)', 'Subcontracts (Rs)'])
        for s in d['stages']:
            sc = d['stage_costs'].get(s.id, {})
            w.writerow([s.name, s.status, f"{s.contract_value:,.2f}",
                        f"{sc.get('labour',0):,.2f}", f"{sc.get('materials',0):,.2f}",
                        f"{sc.get('expenses',0):,.2f}", f"{sc.get('subcontracts',0):,.2f}"])
        w.writerow([])

        # Attendance
        w.writerow(['== Attendance / Labour =='])
        w.writerow(['Date', 'Worker', 'Hours', 'Wage (Rs)', 'Stage', 'Notes'])
        for att, wkr in d['attendance']:
            wage = float(att.wage_calculated or 0)
            stage_name = att.stage.name if att.stage_id else ''
            w.writerow([att.check_in.date(), wkr.name, att.hours, f"{wage:.2f}", stage_name, ''])
        w.writerow([])

        # Expenses
        w.writerow(['== Expenses =='])
        w.writerow(['Date', 'Category', 'Amount (Rs)', 'Stage', 'Remarks'])
        for exp in d['expenses']:
            stage_name = exp.stage.name if exp.stage_id else ''
            w.writerow([exp.date, exp.category, f"{float(exp.amount):,.2f}", stage_name, exp.remarks or ''])
        w.writerow([])

        # Materials
        w.writerow(['== Material Purchases =='])
        w.writerow(['Date', 'Type', 'Material', 'Unit', 'Qty', 'Rate', 'Total (Rs)', 'Stage', 'Supplier', 'Return Ref', 'Reason', 'Approved By'])
        for pur, mat in d['purchases']:
            stage_name = pur.stage.name if pur.stage_id else ''
            w.writerow([pur.date, (pur.entry_type or 'purchase'), mat.name, mat.unit, pur.qty, pur.rate, f"{float(pur.total):,.2f}", stage_name,
                        pur.supplier_name or '', pur.return_ref or '', pur.return_reason or '', pur.approved_by or ''])
        w.writerow([])

        # Subcontract Payments
        w.writerow(['== Subcontract Payments =='])
        w.writerow(['Date', 'Subcontractor', 'Amount (Rs)', 'Stage', 'Notes'])
        for sp, sub in d['sub_payments']:
            stage_name = sub.stage_rel.name if sub.stage_id else ''
            w.writerow([sp.date, sub.name, f"{float(sp.amount):,.2f}", stage_name, sp.notes or ''])
        w.writerow([])

        # Owner Payments
        w.writerow(['== Owner Payments Received =='])
        w.writerow(['Date', 'Amount (Rs)', 'Remarks'])
        for op in d['owner_payments']:
            w.writerow([op.date, f"{float(op.amount):,.2f}", op.remarks or ''])

        out.seek(0)
        fname = f"report_{p.name.replace(' ','_')}.csv"
        return Response(out.getvalue(), mimetype='text/csv',
            headers={'Content-Disposition': f'attachment; filename={fname}'})


    @app.route('/hdc/reports/project/<int:pid>/xlsx')
    @login_required
    def hdc_project_report_xlsx(pid):
        d   = _project_report_data(pid)
        p   = d['project']
        wb  = openpyxl.Workbook()

        # â”€â”€ colour palette
        GREEN  = '87AF32'
        DKGREY = '2D2D2D'
        LTGREY = 'F2F2F2'
        WHITE  = 'FFFFFF'
        RED    = 'C0392B'

        def _hdr(ws, row, cols, label, bg=GREEN, fg=WHITE, bold=True, size=11):
            c = ws.cell(row=row, column=cols[0], value=label)
            c.font       = Font(bold=bold, color=fg, size=size)
            c.fill       = PatternFill('solid', fgColor=bg)
            c.alignment  = Alignment(horizontal='center', vertical='center', wrap_text=True)
            if len(cols) > 1:
                ws.merge_cells(start_row=row, start_column=cols[0],
                               end_row=row, end_column=cols[-1])

        def _row(ws, row, values, bg=None, bold=False, number_format=None):
            for col, val in enumerate(values, 1):
                c = ws.cell(row=row, column=col, value=val)
                c.font      = Font(bold=bold)
                c.alignment = Alignment(vertical='center')
                if bg:
                    c.fill = PatternFill('solid', fgColor=bg)
                if number_format and isinstance(val, (int, float)):
                    c.number_format = number_format

        thin = Side(style='thin', color='CCCCCC')
        def _border(ws, row, ncols):
            for col in range(1, ncols+1):
                ws.cell(row=row, column=col).border = Border(
                    top=thin, bottom=thin, left=thin, right=thin)

        # â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•
        # Sheet 1 â€“ Summary
        # â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•
        ws = wb.active
        ws.title = 'Summary'
        ws.column_dimensions['A'].width = 28
        ws.column_dimensions['B'].width = 22
        ws.row_dimensions[1].height = 30

        _hdr(ws, 1, list(range(1, 3)), 'HDC ERP â€” Project Report', size=14)
        _row(ws, 2, ['Project', p.name], bg=LTGREY, bold=True)
        _row(ws, 3, ['Client', p.client or ''])
        _row(ws, 4, ['Location', p.location or ''])
        _row(ws, 5, ['Contract Type', p.contract_type.replace('_', ' ').title()])
        _row(ws, 6, ['Generated On', _pkt_now().strftime('%Y-%m-%d %H:%M')])
        ws.append([])

        _hdr(ws, 8, list(range(1, 3)), 'Financial Summary', bg=DKGREY)
        fin_rows = [
            ('Contract Value (Rs)', d['contract_value']),
            ('Total Labour Cost (Rs)', d['total_labour']),
            ('Total Material Cost (Rs)', d['total_materials']),
            ('Total Expenses (Rs)', d['total_expenses']),
            ('Total Subcontracts (Rs)', d['total_subcontracts']),
            ('Total Cost (Rs)', d['total_cost']),
            ('Profit / Loss (Rs)', d['profit']),
        ]
        for i, (lbl, val) in enumerate(fin_rows, 9):
            bg = LTGREY if i % 2 == 0 else WHITE
            if lbl.startswith('Total Cost'):
                bg = 'FFE0E0'
            if lbl.startswith('Profit'):
                bg = 'D6EDAF' if val >= 0 else 'FFD5D5'
            _row(ws, i, [lbl, val], bg=bg)
            ws.cell(row=i, column=2).number_format = '#,##0.00'
            _border(ws, i, 2)

        # â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•
        # Sheet 2 â€“ Stage Breakdown
        # â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•
        ws2 = wb.create_sheet('Stage Breakdown')
        hdrs = ['Stage', 'Type', 'Status', 'Stage Value (Rs)',
                'Labour (Rs)', 'Materials (Rs)', 'Expenses (Rs)', 'Subcontracts (Rs)', 'Total Cost (Rs)']
        for i, h in enumerate(hdrs, 1):
            c = ws2.cell(row=1, column=i, value=h)
            c.font  = Font(bold=True, color=WHITE)
            c.fill  = PatternFill('solid', fgColor=GREEN)
            c.alignment = Alignment(horizontal='center')
        ws2.row_dimensions[1].height = 20

        col_w = [28, 14, 12, 18, 16, 16, 16, 18, 18]
        for i, w2 in enumerate(col_w, 1):
            ws2.column_dimensions[get_column_letter(i)].width = w2

        for ri, s in enumerate(d['stages'], 2):
            sc    = d['stage_costs'].get(s.id, {})
            labour    = sc.get('labour', 0)
            mats      = sc.get('materials', 0)
            exp_cost  = sc.get('expenses', 0)
            subc      = sc.get('subcontracts', 0)
            stage_tot = labour + mats + exp_cost + subc
            vals = [s.name, s.contract_basis or '', s.status, s.contract_value,
                    labour, mats, exp_cost, subc, stage_tot]
            bg = LTGREY if ri % 2 == 0 else WHITE
            for ci, val in enumerate(vals, 1):
                c = ws2.cell(row=ri, column=ci, value=val)
                c.fill = PatternFill('solid', fgColor=bg)
                c.alignment = Alignment(vertical='center')
                if ci >= 4:
                    c.number_format = '#,##0.00'
                _border(ws2, ri, len(vals))

        # â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•
        # Sheet 3 â€“ Attendance
        # â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•
        ws3 = wb.create_sheet('Attendance')
        att_hdrs = ['Date', 'Worker', 'Hours', 'Wage (Rs)', 'Stage', 'Notes']
        for i, h in enumerate(att_hdrs, 1):
            c = ws3.cell(row=1, column=i, value=h)
            c.font = Font(bold=True, color=WHITE)
            c.fill = PatternFill('solid', fgColor=GREEN)
        for ci, w3 in enumerate([14, 24, 10, 14, 22, 30], 1):
            ws3.column_dimensions[get_column_letter(ci)].width = w3

        for ri, (att, wkr) in enumerate(d['attendance'], 2):
            wage = float(att.wage_calculated or 0)
            vals = [str(att.check_in.date()), wkr.name, float(att.hours or 0), wage,
                    att.stage.name if att.stage_id else '', '']
            bg = LTGREY if ri % 2 == 0 else WHITE
            for ci, val in enumerate(vals, 1):
                c = ws3.cell(row=ri, column=ci, value=val)
                c.fill = PatternFill('solid', fgColor=bg)
                if ci in (3, 4):
                    c.number_format = '#,##0.00'

        # â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•
        # Sheet 4 â€“ Expenses
        # â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•
        ws4 = wb.create_sheet('Expenses')
        exp_hdrs = ['Date', 'Category', 'Amount (Rs)', 'Stage', 'Description']
        for i, h in enumerate(exp_hdrs, 1):
            c = ws4.cell(row=1, column=i, value=h)
            c.font = Font(bold=True, color=WHITE)
            c.fill = PatternFill('solid', fgColor=GREEN)
        for ci, w4 in enumerate([14, 20, 16, 22, 36], 1):
            ws4.column_dimensions[get_column_letter(ci)].width = w4

        for ri, exp in enumerate(d['expenses'], 2):
            vals = [str(exp.date), exp.category, float(exp.amount or 0),
                    exp.stage.name if exp.stage_id else '', exp.remarks or '']
            bg = LTGREY if ri % 2 == 0 else WHITE
            for ci, val in enumerate(vals, 1):
                c = ws4.cell(row=ri, column=ci, value=val)
                c.fill = PatternFill('solid', fgColor=bg)
                if ci == 3:
                    c.number_format = '#,##0.00'

        # â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•
        # Sheet 5 â€“ Materials
        # â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•
        ws5 = wb.create_sheet('Materials')
        mat_hdrs = ['Date', 'Type', 'Material', 'Unit', 'Qty', 'Rate (Rs)', 'Total (Rs)', 'Stage', 'Supplier', 'Return Ref', 'Reason', 'Approved By']
        for i, h in enumerate(mat_hdrs, 1):
            c = ws5.cell(row=1, column=i, value=h)
            c.font = Font(bold=True, color=WHITE)
            c.fill = PatternFill('solid', fgColor=GREEN)
        for ci, w5 in enumerate([14, 12, 24, 10, 10, 14, 16, 22, 18, 14, 20, 18], 1):
            ws5.column_dimensions[get_column_letter(ci)].width = w5

        for ri, (pur, mat) in enumerate(d['purchases'], 2):
            vals = [str(pur.date), (pur.entry_type or 'purchase'), mat.name, mat.unit, float(pur.qty or 0),
                    float(pur.rate or 0), float(pur.total or 0),
                    pur.stage.name if pur.stage_id else '',
                    pur.supplier_name or '',
                    pur.return_ref or '',
                    pur.return_reason or '',
                    pur.approved_by or '']
            bg = LTGREY if ri % 2 == 0 else WHITE
            for ci, val in enumerate(vals, 1):
                c = ws5.cell(row=ri, column=ci, value=val)
                c.fill = PatternFill('solid', fgColor=bg)
                if ci in (5, 6, 7):
                    c.number_format = '#,##0.00'

        # â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•
        # Sheet 6 â€“ Subcontract Payments
        # â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•
        ws6 = wb.create_sheet('Subcontract Payments')
        sp_hdrs = ['Date', 'Subcontractor', 'Amount (Rs)', 'Stage', 'Notes']
        for i, h in enumerate(sp_hdrs, 1):
            c = ws6.cell(row=1, column=i, value=h)
            c.font = Font(bold=True, color=WHITE)
            c.fill = PatternFill('solid', fgColor=GREEN)
        for ci, w6 in enumerate([14, 28, 16, 22, 34], 1):
            ws6.column_dimensions[get_column_letter(ci)].width = w6

        for ri, (sp, sub) in enumerate(d['sub_payments'], 2):
            vals = [str(sp.date), sub.name, float(sp.amount or 0),
                    sub.stage_rel.name if sub.stage_id else '', sp.notes or '']
            bg = LTGREY if ri % 2 == 0 else WHITE
            for ci, val in enumerate(vals, 1):
                c = ws6.cell(row=ri, column=ci, value=val)
                c.fill = PatternFill('solid', fgColor=bg)
                if ci == 3:
                    c.number_format = '#,##0.00'

        # â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•
        # Sheet 7 â€“ Owner Payments
        # â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•
        ws7 = wb.create_sheet('Owner Payments')
        op_hdrs = ['Date', 'Amount (Rs)', 'Remarks']
        for i, h in enumerate(op_hdrs, 1):
            c = ws7.cell(row=1, column=i, value=h)
            c.font = Font(bold=True, color=WHITE)
            c.fill = PatternFill('solid', fgColor=GREEN)
        for ci, w7 in enumerate([14, 16, 40], 1):
            ws7.column_dimensions[get_column_letter(ci)].width = w7

        for ri, op in enumerate(d['owner_payments'], 2):
            vals = [str(op.date), float(op.amount or 0), op.remarks or '']
            bg = LTGREY if ri % 2 == 0 else WHITE
            for ci, val in enumerate(vals, 1):
                c = ws7.cell(row=ri, column=ci, value=val)
                c.fill = PatternFill('solid', fgColor=bg)
                if ci == 2:
                    c.number_format = '#,##0.00'

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        fname = f"report_{p.name.replace(' ','_')}.xlsx"
        return Response(buf.getvalue(),
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={'Content-Disposition': f'attachment; filename={fname}'})


    @app.route('/hdc/reports/project/<int:pid>/pdf')
    @login_required
    def hdc_project_report_pdf(pid):
        d = _project_report_data(pid)
        return render_template('projects/project_report_print.html', **d,
                               now=_pkt_now().strftime('%Y-%m-%d %H:%M'))
